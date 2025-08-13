"""Loading and rendering prompt templates."""
from __future__ import annotations

import json
from pathlib import Path
from typing import Dict, Iterable, List, Sequence, Union

from .errorlog import get_logger

_log = get_logger(__name__)


def read_file_safe(path: str) -> str:
    """Return file contents or empty string with logging."""
    p = Path(path).expanduser()
    try:
        if p.suffix.lower() == ".docx":
            try:
                import docx  # type: ignore
                return "\n".join(par.text for par in docx.Document(p).paragraphs)
            except Exception as e:  # pragma: no cover - optional dependency
                _log.error("cannot read Word file %s: %s", path, e)
                return ""
        return p.read_text()
    except Exception as e:
        _log.error("cannot read file %s: %s", path, e)
        return ""


def load_template(path: Path) -> Dict:
    """Load JSON template file with validation."""
    path = path.expanduser().resolve()
    if not path.is_file():
        _log.error("template not found: %s", path)
        raise FileNotFoundError(path)
    with path.open(encoding="utf-8") as fh:
        return json.load(fh)


def validate_template(data: Dict) -> bool:
    """Basic schema validation."""
    required = {"id", "title", "style", "template", "placeholders"}
    return required.issubset(data)


def fill_placeholders(
    lines: Iterable[str], vars: Dict[str, Union[str, Sequence[str], None]]
) -> str:
    """Replace ``{{name}}`` placeholders with values.

    ``vars`` may contain either strings or sequences of strings. Sequence
    values (used for dynamic list placeholders) are joined with newlines before
    replacement. ``None`` or empty string values cause the entire line containing
    the placeholder to be removed.
    """

    out: List[str] = []
    for line in lines:
        skip_line = False
        for k, v in vars.items():
            placeholder = f"{{{{{k}}}}}"
            if placeholder not in line:
                continue

            if v is None:
                skip_line = True
                break
            if isinstance(v, (list, tuple)):
                repl = "\n".join(str(item) for item in v)
                if not repl.strip():
                    skip_line = True
                    break
            else:
                repl = str(v)
                if not repl.strip():
                    skip_line = True
                    break
            line = line.replace(placeholder, repl)

        if not skip_line:
            out.append(line)
    return "\n".join(out)
